"""Pruebas unitarias para el módulo de factores de precios y autocorrección de líneas."""

from __future__ import annotations

from pathlib import Path

import pytest

from app.services.almacen.factores import (
    CATALOGO_FACTORES,
    CATALOGO_LINEAS_BASE,
    SIN_FACTOR,
    calcular_precios_numericos,
    cargar_factores_personalizados,
    detectar_rollo_metros,
    formatear_menu_factores,
    generar_formulas_fila,
    guardar_factor_personalizado,
    normalizar_eleccion_factor,
    recuperar_clave_truncada,
    resolver_linea,
)


def test_catalogo_base_tiene_8_factores_y_76_lineas() -> None:
    assert len(CATALOGO_FACTORES) == 8
    assert len(CATALOGO_LINEAS_BASE) == 76
    assert CATALOGO_LINEAS_BASE["FESTO"]["factor"] == "FESTO"
    assert CATALOGO_LINEAS_BASE["CYB"]["factor"] == "CYBERMATIC"
    assert CATALOGO_LINEAS_BASE["DAYCO"]["factor"] is None


@pytest.mark.parametrize(
    ("linea_raw", "proveedor", "descripcion", "esperada", "factor_esperado"),
    [
        ("FEST", "FESTO PNEUMATIC SA", "CILINDRO", "FESTO", "FESTO"),
        ("ADC", "ADCCOM", "CONECTOR", "ADCCO", "COMPRAS GENERALES FORANEAS"),
        ("GEN", "GENEBRE", "VALVULA", "GENEB", "COMPRAS GENERALES FORANEAS"),
        ("KRUIY", "KURIYAMA", "MANGUERA", "KURIY", None),
        ("PKR", "PARKER FLUID", "MANGUERA HIDRAULICA", "PKRFC", "COMPRAS GENERALES FORANEAS"),
        ("PKR", "PARKER", "ELEMENTO FILTRANTE", "PKRF", "COMPRAS GENERALES FORANEAS"),
        ("VAR", "PROVEEDOR", "CABLE ELECTRICO", "VAREL", "COMPRAS GENERALES FORANEAS"),
        ("MAN", "DEWIT", "MANOMETRO GLICERINA 100 PSI", "MANOM", None),
        ("VAY", "NUEVO PROV", "PRODUCTO NUEVO", "VAY", None),
    ],
)
def test_resolver_linea_autocorrige_codigos_truncados(
    linea_raw: str,
    proveedor: str,
    descripcion: str,
    esperada: str,
    factor_esperado: str | None,
    tmp_path: Path,
) -> None:
    ruta_json = tmp_path / "FACTORES_PERSONALIZADOS.json"
    resultado = resolver_linea(
        linea_raw,
        proveedor=proveedor,
        descripcion=descripcion,
        ruta_personalizados=ruta_json,
    )
    assert resultado["linea_resuelta"] == esperada
    assert resultado["factor"] == factor_esperado


def test_guardar_y_cargar_factor_personalizado(tmp_path: Path) -> None:
    ruta_json = tmp_path / "FACTORES_PERSONALIZADOS.json"
    guardar_factor_personalizado(
        linea="DAYCO",
        factor="2",
        descripcion="DAYCO",
        ruta=ruta_json,
    )
    guardar_factor_personalizado(
        linea="VAY",
        factor="COMPRAS GENERALES LOCALES",
        descripcion="LINEA VAY",
        ruta=ruta_json,
    )

    guardados = cargar_factores_personalizados(ruta_json)
    assert guardados["DAYCO"]["factor"] == "COMPRAS GENERALES FORANEAS"
    assert guardados["VAY"]["factor"] == "COMPRAS GENERALES LOCALES"

    resuelto = resolver_linea("DAY", ruta_personalizados=ruta_json)
    assert resuelto["linea_resuelta"] == "DAYCO"
    assert resuelto["factor"] == "COMPRAS GENERALES FORANEAS"


def test_normalizar_eleccion_factor_y_menu() -> None:
    menu = formatear_menu_factores()
    assert "1. **COMPRAS GENERALES LOCALES**" in menu
    assert "8. **CYBERMATIC**" in menu
    assert "9. **Sin factor**" in menu

    assert normalizar_eleccion_factor(1) == "COMPRAS GENERALES LOCALES"
    assert normalizar_eleccion_factor("2") == "COMPRAS GENERALES FORANEAS"
    assert normalizar_eleccion_factor("8. CYBERMATIC") == "CYBERMATIC"
    assert normalizar_eleccion_factor("9") == SIN_FACTOR
    assert normalizar_eleccion_factor("sin factor") == SIN_FACTOR


def test_recuperar_clave_truncada_en_txt() -> None:
    assert recuperar_clave_truncada("PU14TRANSBLU", "PU14TRANSBLUE") == "PU14TRANSBLUE"
    assert recuperar_clave_truncada("65040806", "6504 08 06") == "65040806"
    assert recuperar_clave_truncada("", "ABC123") == "ABC123"


def test_detectar_rollo_metros() -> None:
    caso_cyber = detectar_rollo_metros(
        clave_unidad="XRO",
        descripcion_xml="(100m) Manguera de Poliuretano en Carret",
        descripcion_apoyo="MANGUERA POLIURETANO 10X6.5 AZUL TR",
    )
    assert caso_cyber["es_xro"] is True
    assert caso_cyber["metros"] == 100.0
    assert caso_cyber["requiere_confirmacion"] is True

    caso_pieza_normal = detectar_rollo_metros(
        clave_unidad="H87",
        descripcion_xml="Conector Recto Macho 10MM x 1/4 NPT",
        descripcion_apoyo="CONECTOR RECTO",
    )
    assert caso_pieza_normal["es_xro"] is False
    assert caso_pieza_normal["metros"] is None
    assert caso_pieza_normal["requiere_confirmacion"] is False


def test_calcular_precios_numericos_y_formulas_cybermatic() -> None:
    precios = calcular_precios_numericos("CYBERMATIC", 90.0)
    assert precios == {
        "PUBLICO": 200.0,
        "MINIMO": 160.0,
        "MOSTRADOR": 200.0,
        "MAYOREO": 180.0,
        "DISTRIBUIDOR": 160.0,
    }

    formulas = generar_formulas_fila(
        factor_nombre="CYBERMATIC",
        fila_excel=6,
        precio_compra_bruto=1796.3236,
        divisor_metros=100.0,
    )
    assert formulas == {
        "PRECIO COMPRA": "=1796.3236/100",
        "PUBLICO": "=S6/.9*2",
        "MINIMO": "=T6*.8",
        "MOSTRADOR": "=T6",
        "MAYOREO": "=T6*.9",
        "DISTRIBUIDOR": "=T6*.8",
    }


def test_excel_inyecta_formulas_dinamicas_y_txt_exporta_numeros(tmp_path: Path) -> None:
    import openpyxl
    import pandas as pd
    from app.services.almacen import excel

    concepto_rollo = {
        "no_identificacion": "PU14TRANSBLUE",
        "codigo_secundario": None,
        "descripcion": "(100m) Manguera de Poliuretano en Carret",
        "clave_prod_serv": "40142000",
        "clave_unidad": "XRO",
        "cantidad": 1.0,
        "valor_unitario": 1796.3236,
        "descuento": 0.0,
        "proveedor_nombre": "CYBERMATICS",
        "fecha_compra": "2026-08-04T12:00:00",
    }
    dato_txt_rollo = {
        "clave_articulo": "PU14TRANSBLU",  # truncado a 12 chars por SAE
        "linea": "CYB",
        "descripcion_corta": "MANGUERA POLIURETANO 10X6.5 AZUL TR",
        "numero_proveedor": "491",
    }

    meta_por_clave: dict[str, dict[str, object]] = {}
    fila = excel._fila_desde_concepto(
        concepto_rollo,
        dato_txt_rollo,
        decisiones_rollos={
            "PU14TRANSBLUE": {
                "cambiar_a_mtr": True,
                "dividir_entre_metros": True,
                "metros": 100.0,
            }
        },
        meta_por_clave=meta_por_clave,
    )

    assert fila["Clave Artículo"] == "PU14TRANSBLUE"
    assert fila["Clave unidad"] == "MTR"
    assert fila["Línea"] == "CYB"
    assert fila["PRECIO COMPRA"] == 17.9632
    assert fila["PUBLICO"] == round(17.9632 / 0.9 * 2, 4)

    df = pd.DataFrame([fila], columns=excel.COLUMNAS_ASPEL)
    df.attrs["meta_por_clave"] = meta_por_clave

    ruta_xlsx = excel.guardar_en_base_acumulada(df, str(tmp_path), prefijo="TEST_REPORTE")

    # 1. Verificar fórmulas dinámicas en el .xlsx
    wb = openpyxl.load_workbook(ruta_xlsx, data_only=False)
    ws = wb.active
    assert ws is not None
    encabezados = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    assert ws.cell(row=2, column=encabezados["PRECIO COMPRA"]).value == "=1796.3236/100"
    assert ws.cell(row=2, column=encabezados["PUBLICO"]).value == "=S2/.9*2"
    assert ws.cell(row=2, column=encabezados["MINIMO"]).value == "=T2*.8"
    assert ws.cell(row=2, column=encabezados["MOSTRADOR"]).value == "=T2"
    assert ws.cell(row=2, column=encabezados["MAYOREO"]).value == "=T2*.9"
    assert ws.cell(row=2, column=encabezados["DISTRIBUIDOR"]).value == "=T2*.8"

    # 1b. Verificar valores en caché (<v>) para que Excel los muestre en Vista Protegida
    #     sin necesidad de pulsar "Habilitar edición"
    wb_cache = openpyxl.load_workbook(ruta_xlsx, data_only=True)
    ws_cache = wb_cache.active
    assert ws_cache is not None
    assert ws_cache.cell(row=2, column=encabezados["PRECIO COMPRA"]).value == 17.9632
    assert ws_cache.cell(row=2, column=encabezados["PUBLICO"]).value == 39.9182
    assert ws_cache.cell(row=2, column=encabezados["MINIMO"]).value == 31.9346
    assert ws_cache.cell(row=2, column=encabezados["MOSTRADOR"]).value == 39.9182
    assert ws_cache.cell(row=2, column=encabezados["MAYOREO"]).value == 35.9264
    assert ws_cache.cell(row=2, column=encabezados["DISTRIBUIDOR"]).value == 31.9346

    # 2. Verificar valores numéricos en el TXT (sin fórmulas '=' ni NaN)
    ruta_txt_comas = tmp_path / "TEST_REPORTE_COMAS.txt"
    contenido_txt = ruta_txt_comas.read_text(encoding="utf-8")
    assert "=S2" not in contenido_txt
    assert "nan" not in contenido_txt.lower()
    assert "17.9632" in contenido_txt
    assert "39.9182" in contenido_txt


def _xml_cfdi_ejemplo(
    *,
    no_identificacion: str = "PROD01",
    descripcion: str = "CONECTOR RECTO 1/4",
    clave_unidad: str = "H87",
    valor_unitario: str = "100.00",
    proveedor: str = "PROVEEDOR DEMO SA",
    folio: str = "100",
) -> bytes:
    return f"""<?xml version="1.0" encoding="utf-8"?>
<cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/4" Folio="{folio}" Fecha="2026-08-04T12:00:00">
  <cfdi:Emisor Nombre="{proveedor}" Rfc="PDE010101AAA"/>
  <cfdi:Conceptos>
    <cfdi:Concepto
      NoIdentificacion="{no_identificacion}"
      Descripcion="{descripcion}"
      ClaveProdServ="40142000"
      ClaveUnidad="{clave_unidad}"
      Cantidad="1"
      ValorUnitario="{valor_unitario}"
      Importe="{valor_unitario}"/>
  </cfdi:Conceptos>
</cfdi:Comprobante>""".encode("utf-8")


def _txt_sae_ejemplo(
    *,
    clave: str = "PROD01",
    descripcion_corta: str = "CONECTOR RECTO 1/4",
    linea: str = "FESTO",
    folio: str = "100",
) -> bytes:
    return f"""COMPRA  {folio}
Proveedor: ( 491 ) PROVEEDOR DEMO SA
  1.00  {clave}  {linea}  {descripcion_corta}
""".encode("utf-8")


def _configurar_drive_mock(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    carpeta_nombre: str,
    archivos_contenido: dict[str, bytes],
) -> None:
    from app.agents import tools
    from app.core.config import settings
    from app.services.almacen import excel
    from app.services.almacen.fuentes import google_drive_client, sincronizador
    from app.services.almacen.fuentes.google_drive import CarpetaFactura, DriveArchivo

    monkeypatch.setattr(settings, "google_drive_folder_id", "docs-root")
    monkeypatch.setattr(settings, "google_service_account_json", '{"type": "service_account"}')
    monkeypatch.setattr(tools, "CARPETA_DATOS", tmp_path)
    monkeypatch.setattr(sincronizador, "CARPETA_DATOS", tmp_path)
    monkeypatch.setattr(excel, "CARPETA_DATOS", tmp_path)

    monkeypatch.setattr(sincronizador, "crear_cliente_drive", lambda: object())
    monkeypatch.setattr(
        sincronizador,
        "reubicar_carpetas_mal_ubicadas",
        lambda c, cid: {"carpetas_reubicadas": [], "archivos_reubicados_ids": []},
    )
    monkeypatch.setattr(sincronizador, "buscar_carpeta_hija", lambda *a, **kw: "input-root")
    monkeypatch.setattr(sincronizador, "buscar_o_crear_carpeta", lambda c, nombre, padre: f"id-{nombre}")

    drive_archivos = tuple(
        DriveArchivo(
            f"id-{nombre}",
            nombre,
            "text/xml" if nombre.endswith(".xml") else "text/plain",
            "v1",
            "sub-1",
        )
        for nombre in archivos_contenido
    )
    carpeta = CarpetaFactura(id="sub-1", nombre=carpeta_nombre, archivos=drive_archivos)
    monkeypatch.setattr(sincronizador, "listar_carpetas_factura", lambda c, cid: [carpeta])
    monkeypatch.setattr(
        sincronizador,
        "descargar_archivo",
        lambda c, arch, destino: destino.write_bytes(archivos_contenido[arch.nombre]),
    )
    monkeypatch.setattr(google_drive_client, "mover_archivo", lambda *a, **kw: None)
    monkeypatch.setattr(google_drive_client, "subir_archivo", lambda *a, **kw: None)
    monkeypatch.setattr(google_drive_client, "eliminar_archivo", lambda *a, **kw: None)


def test_confirmacion_interactiva_carpeta_solo_xml(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    import json
    from app.agents import tools

    _configurar_drive_mock(
        monkeypatch,
        tmp_path,
        "ENTRADA_SOLO_XML",
        {"factura-100.xml": _xml_cfdi_ejemplo()},
    )

    # 1. Sin confirmación: debe detenerse y preguntar y/n
    resp_1 = json.loads(tools.generar_entradas_almacen.invoke({}))
    assert resp_1["ok"] is True
    assert resp_1["requiere_confirmacion"] is True
    assert resp_1["tipo_confirmacion"] == "solo_xml"
    assert "ENTRADA_SOLO_XML" in resp_1["pregunta_usuario"]
    assert "(**y/n**)" in resp_1["pregunta_usuario"]

    # 2. Usuario responde 'n' (confirmar_solo_xml=False): cancela sin generar reporte
    resp_no = json.loads(tools.generar_entradas_almacen.invoke({"confirmar_solo_xml": False}))
    assert resp_no["ok"] is True
    assert resp_no["reporte_generado"] is False
    assert resp_no["cancelado_por_usuario"] is True
    assert "revisar que hayas subido todos los documentos" in resp_no["mensaje"]

    # 3. Usuario responde 'y' (confirmar_solo_xml=True): genera el reporte con éxito
    resp_si = json.loads(tools.generar_entradas_almacen.invoke({"confirmar_solo_xml": True}))
    assert resp_si["ok"] is True
    assert resp_si["reporte_generado"] is True
    assert resp_si["resumen"]["productos_nuevos"] == 1


def test_confirmacion_interactiva_linea_sin_factor_y_guardado(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    import json
    from app.agents import tools
    from app.services.almacen import factores

    ruta_json_custom = tmp_path / "FACTORES_PERSONALIZADOS.json"
    monkeypatch.setattr(factores, "RUTA_FACTORES_PERSONALIZADOS", ruta_json_custom)

    _configurar_drive_mock(
        monkeypatch,
        tmp_path,
        "ENTRADA_DAYCO",
        {
            "factura-100.xml": _xml_cfdi_ejemplo(no_identificacion="BANDA01", valor_unitario="200.00"),
            "entrada-100.txt": _txt_sae_ejemplo(clave="BANDA01", linea="DAYCO"),
        },
    )

    # 1. Primera llamada: detecta que DAYCO no tiene factor definido y muestra menú 1..9
    resp_1 = json.loads(tools.generar_entradas_almacen.invoke({}))
    assert resp_1["requiere_confirmacion"] is True
    assert resp_1["tipo_confirmacion"] == "lineas_sin_factor"
    assert "DAYCO" in resp_1["pregunta_usuario"]
    assert "1. **COMPRAS GENERALES LOCALES**" in resp_1["pregunta_usuario"]
    assert "¿Deseas guardar esta respuesta" in resp_1["pregunta_usuario"]

    # 2. Usuario elige factor 2 (COMPRAS GENERALES FORANEAS) y guardar_factores=True
    resp_2 = json.loads(
        tools.generar_entradas_almacen.invoke(
            {
                "factores_elegidos": {"DAYCO": "2"},
                "guardar_factores": True,
            }
        )
    )
    assert resp_2["ok"] is True
    assert resp_2["reporte_generado"] is True
    assert resp_2["resumen"]["lineas"] == ["DAYCO"]

    guardados = factores.cargar_factores_personalizados(ruta_json_custom)
    assert guardados["DAYCO"]["factor"] == "COMPRAS GENERALES FORANEAS"


def test_confirmacion_interactiva_rollos_xro_y_metros(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    import json
    import pandas as pd
    from app.agents import tools

    _configurar_drive_mock(
        monkeypatch,
        tmp_path,
        "ENTRADA_CYB",
        {
            "cyber-100.xml": _xml_cfdi_ejemplo(
                no_identificacion="PU14TRANSBLUE",
                descripcion="(100m) Manguera de Poliuretano en Carret",
                clave_unidad="XRO",
                valor_unitario="1796.3236",
                proveedor="CYBERMATICS",
            ),
            "cyber-100.txt": _txt_sae_ejemplo(
                clave="PU14TRANSBLU",
                descripcion_corta="MANGUERA POLIURETANO 10X6.5 AZUL TR",
                linea="CYB",
            ),
        },
    )

    # 1. Primera llamada: detecta rollo XRO con (100m) y pide confirmación y/n
    resp_1 = json.loads(tools.generar_entradas_almacen.invoke({}))
    assert resp_1["requiere_confirmacion"] is True
    assert resp_1["tipo_confirmacion"] == "rollos_pendientes"
    assert "PU14TRANSBLUE" in resp_1["pregunta_usuario"]
    assert "MTR" in resp_1["pregunta_usuario"]
    assert "100" in resp_1["pregunta_usuario"]

    # 2. Usuario confirma ambos (cambiar_xro_a_mtr=True, dividir_precio_entre_metros=True)
    resp_2 = json.loads(
        tools.generar_entradas_almacen.invoke(
            {
                "cambiar_xro_a_mtr": True,
                "dividir_precio_entre_metros": True,
            }
        )
    )
    assert resp_2["ok"] is True
    assert resp_2["reporte_generado"] is True

    df_generado = pd.read_excel(tmp_path / "BASE_ENTRADAS_ALMACEN.xlsx")
    assert df_generado.iloc[0]["Clave Artículo"] == "PU14TRANSBLUE"
    assert df_generado.iloc[0]["Clave unidad"] == "MTR"


def test_actualizar_precios_reporte_existente_aplica_precios_y_formulas(
    tmp_path: Path,
) -> None:
    import openpyxl
    import pandas as pd
    from app.services.almacen import excel

    # Crear un reporte Excel ya generado sin precios calculados (en None / 0) y con línea truncada
    fila_sin_precios = {col: None for col in excel.COLUMNAS_ASPEL}
    fila_sin_precios.update(
        {
            "Clave Artículo": "PROD01",
            "DESCRIPCION": "CILINDRO NEUMATICO",
            "Línea": "FEST",
            "Clave unidad": "H87",
            "PROVEEDOR": "FESTO PNEUMATIC SA",
            "PRECIO COMPRA": 100.0,
            "PUBLICO": 0.0,
            "MINIMO": None,
            "MOSTRADOR": None,
            "MAYOREO": None,
            "DISTRIBUIDOR": None,
        }
    )
    ruta_xlsx = tmp_path / "ENTRADAS_ALMACEN_13-04-2026.xlsx"
    pd.DataFrame([fila_sin_precios], columns=excel.COLUMNAS_ASPEL).to_excel(
        ruta_xlsx, index=False
    )

    resultado = excel.actualizar_precios_reporte_existente(str(ruta_xlsx))
    assert resultado["actualizado"] is True
    assert resultado["filas_actualizadas"] == 1

    # Verificar que la línea se autocorrigió y los TXT tienen precios numéricos
    ruta_txt = tmp_path / "ENTRADAS_ALMACEN_13-04-2026_COMAS.txt"
    assert ruta_txt.is_file()
    contenido_txt = ruta_txt.read_text(encoding="utf-8")
    assert "FESTO" in contenido_txt
    assert "142.8571" in contenido_txt  # 100 / 0.7

    # Verificar que el Excel contiene las fórmulas dinámicas
    wb = openpyxl.load_workbook(ruta_xlsx, data_only=False)
    ws = wb.active
    assert ws is not None
    encabezados = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    assert ws.cell(row=2, column=encabezados["Línea"]).value == "FESTO"
    assert ws.cell(row=2, column=encabezados["PUBLICO"]).value == "=S2/.7"
    assert ws.cell(row=2, column=encabezados["MINIMO"]).value == "=S2/.9"


def test_actualizar_precios_reportes_drive_actualiza_reportes_historicos_y_base(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    import json
    import openpyxl
    import pandas as pd
    from app.agents import tools
    from app.core.config import settings
    from app.services.almacen import excel
    from app.services.almacen.fuentes import google_drive_client, sincronizador
    from app.services.almacen.fuentes.google_drive import MIME_CARPETA

    monkeypatch.setattr(settings, "google_drive_folder_id", "root-drive")
    monkeypatch.setattr(tools, "CARPETA_DATOS", tmp_path)
    monkeypatch.setattr(sincronizador, "CARPETA_DATOS", tmp_path)
    monkeypatch.setattr(excel, "CARPETA_DATOS", tmp_path)
    monkeypatch.setattr(sincronizador, "crear_cliente_drive", lambda: object())

    # Crear un Excel sin precios ni línea para simular reporte viejo en Drive y en base local
    fila_vieja = {col: None for col in excel.COLUMNAS_ASPEL}
    fila_vieja.update(
        {
            "Clave Artículo": "PROD01",
            "DESCRIPCION": "CONECTOR RECTO 1/4",
            "Línea": None,
            "Clave unidad": "H87",
            "PROVEEDOR": "PROVEEDOR DEMO SA",
            "PRECIO COMPRA": 100.0,
        }
    )
    ruta_base_local = tmp_path / "BASE_ENTRADAS_ALMACEN.xlsx"
    pd.DataFrame([fila_vieja], columns=excel.COLUMNAS_ASPEL).to_excel(
        ruta_base_local, index=False
    )
    bytes_xlsx_sin_precios = ruta_base_local.read_bytes()

    hijos_por_carpeta = {
        "root-drive": [
            {"id": "doc-entrada", "name": "DOCUMENTOS ENTRADA", "mimeType": MIME_CARPETA}
        ],
        "doc-entrada": [
            {"id": "fecha-1", "name": "13-04-2026", "mimeType": MIME_CARPETA}
        ],
        "fecha-1": [
            {"id": "prov-1", "name": "Proveedor Demo - 13-04-2026", "mimeType": MIME_CARPETA},
            {
                "id": "xlsx-batch",
                "name": "ENTRADAS_ALMACEN_13-04-2026.xlsx",
                "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            },
        ],
        "prov-1": [
            {"id": "xml-1", "name": "factura-100.xml", "mimeType": "text/xml"},
            {"id": "txt-1", "name": "entrada-100.txt", "mimeType": "text/plain"},
            {
                "id": "xlsx-prov",
                "name": "ENTRADAS_ALMACEN_Proveedor_Demo_13-04-2026.xlsx",
                "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            },
        ],
    }

    contenido_drive = {
        "factura-100.xml": _xml_cfdi_ejemplo(),
        "entrada-100.txt": _txt_sae_ejemplo(linea="FESTO"),
        "ENTRADAS_ALMACEN_13-04-2026.xlsx": bytes_xlsx_sin_precios,
        "ENTRADAS_ALMACEN_Proveedor_Demo_13-04-2026.xlsx": bytes_xlsx_sin_precios,
    }
    subidos: list[tuple[str, str]] = []

    monkeypatch.setattr(
        google_drive_client,
        "_listar_hijos",
        lambda _c, cid: hijos_por_carpeta.get(cid, []),
    )
    monkeypatch.setattr(
        sincronizador,
        "descargar_archivo",
        lambda _c, arch, destino: destino.write_bytes(contenido_drive[arch.nombre]),
    )
    monkeypatch.setattr(google_drive_client, "eliminar_archivo", lambda *a, **kw: None)
    monkeypatch.setattr(
        google_drive_client,
        "subir_archivo",
        lambda _c, ruta, cid: subidos.append((ruta.name, cid)),
    )

    resp = json.loads(tools.actualizar_precios_reportes_generados.invoke({}))
    assert resp["ok"] is True
    assert resp["reportes_revisados"] == 3
    assert resp["total_actualizados"] == 3
    assert resp["base_local_actualizada"] is True

    # Verificar que la base local recuperó la línea FESTO y tiene las fórmulas
    wb = openpyxl.load_workbook(ruta_base_local, data_only=False)
    ws = wb.active
    assert ws is not None
    encabezados = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    assert ws.cell(row=2, column=encabezados["Línea"]).value == "FESTO"
    assert ws.cell(row=2, column=encabezados["PUBLICO"]).value == "=S2/.7"

